from pathlib import Path
import sys, time, json, warnings
ROOT = Path(__file__).resolve().parent
PKG = ROOT.parent / 'backend' / '.packages'
sys.path.insert(0, str(PKG))
import numpy as np
import pandas as pd
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from xgboost import XGBRegressor
import joblib
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

SEED = 42
DATA = ROOT.parent / 'dataset'
TRAIN_FILE = DATA / 'train.csv'
TEST_FILE = DATA / 'test.csv'
OUT_XLSX = ROOT / 'model_experiment_results.xlsx'
OUT_MODEL = ROOT / 'best_model.pkl'
TARGET = 'accident_risk'
DROP = {'id', TARGET}


def make_pipeline(X, estimator):
    cat = X.select_dtypes(include=['object', 'category']).columns.tolist()
    numeric = [c for c in X.columns if c not in cat]
    prep = ColumnTransformer([
        ('numeric', Pipeline([('imputer', SimpleImputer(strategy='median')), ('scale', StandardScaler())]), numeric),
        ('categorical', Pipeline([('imputer', SimpleImputer(strategy='most_frequent')), ('onehot', OneHotEncoder(handle_unknown='ignore'))]), cat),
    ], remainder='drop')
    return Pipeline([('preprocess', prep), ('model', estimator)])


def algorithms():
    return [
        ('Linear Regression', LinearRegression()),
        ('Random Forest Regressor', RandomForestRegressor(n_estimators=200, min_samples_leaf=2, n_jobs=-1, random_state=SEED)),
        ('XGBoost Regressor', XGBRegressor(n_estimators=250, max_depth=6, learning_rate=0.08, subsample=0.85,
             colsample_bytree=0.85, reg_lambda=1.0, objective='reg:squarederror', tree_method='hist',
             n_jobs=4, random_state=SEED, verbosity=0)),
    ]


def main():
    if not TRAIN_FILE.exists() or not TEST_FILE.exists():
        raise FileNotFoundError(f'Expected {TRAIN_FILE} and {TEST_FILE}')
    frame = pd.read_csv(TRAIN_FILE)
    if TARGET not in frame:
        raise ValueError(f'Missing target column: {TARGET}')
    X_all = frame.drop(columns=[c for c in DROP if c in frame.columns])
    y_all = frame[TARGET].astype(float)
    X_train, X_val, y_train, y_val = train_test_split(X_all, y_all, test_size=0.2, random_state=SEED)
    rows, candidates = [], []
    for include_accidents in (True, False):
        Xtr = X_train.copy()
        Xva = X_val.copy()
        if not include_accidents:
            Xtr = Xtr.drop(columns=['num_reported_accidents'])
            Xva = Xva.drop(columns=['num_reported_accidents'])
        for name, estimator in algorithms():
            pipe = make_pipeline(Xtr, estimator)
            start = time.perf_counter()
            pipe.fit(Xtr, y_train)
            seconds = time.perf_counter() - start
            pred = pipe.predict(Xva)
            metrics = {'Algorithm': name, 'Accident Count Included': 'Yes' if include_accidents else 'No',
                       'MAE': mean_absolute_error(y_val, pred), 'RMSE': float(np.sqrt(mean_squared_error(y_val, pred))),
                       'R²': r2_score(y_val, pred), 'Training Time (seconds)': seconds}
            rows.append(metrics)
            candidates.append((metrics, include_accidents, name, estimator))
            print(f"Completed {name} | accident count={'yes' if include_accidents else 'no'} | MAE={metrics['MAE']:.5f} | RMSE={metrics['RMSE']:.5f} | R2={metrics['R²']:.5f} | seconds={seconds:.1f}", flush=True)
    # Select strictly on validation MAE, with RMSE as the tie-breaker.
    winner, use_count, winning_name, winning_estimator = min(candidates, key=lambda c: (c[0]['MAE'], c[0]['RMSE']))
    result = pd.DataFrame(rows)
    with pd.ExcelWriter(OUT_XLSX, engine='openpyxl') as writer:
        result.to_excel(writer, index=False, sheet_name='Validation Results')
        pd.DataFrame([
            ['Validation split', '80/20 random split; random_state=42; shared by all six runs'],
            ['Selection rule', 'Lowest validation MAE; RMSE breaks ties. Test data was not used for selection.'],
            ['Best validation run', f"{winning_name} | accident count {'included' if use_count else 'excluded'}"],
            ['Leakage note', 'num_reported_accidents may be unavailable at prospective prediction time; compare both configurations and decide from the intended prediction timestamp.'],
            ['Test evaluation', 'Unavailable: test.csv has no accident_risk labels.'],
            ['Saved model', 'best_model.pkl is refit on all labeled training rows using the selected validation configuration.'],
            ['Random seed', SEED],
        ], columns=['Item', 'Details']).to_excel(writer, index=False, sheet_name='Notes')
    wb = load_workbook(OUT_XLSX)
    ws = wb['Validation Results']
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78'); cell.alignment = Alignment(wrap_text=True)
    for row in range(2, ws.max_row + 1):
        cfg = ws.cell(row, 2).value
        fill = 'EAF2F8' if cfg == 'Yes' else 'F4F6F7'
        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).fill = PatternFill('solid', fgColor=fill)
        for col in (3,4,5,6): ws.cell(row,col).number_format = '0.0000'
    best_row = 2 + rows.index(winner)
    for cell in ws[best_row]:
        cell.fill = PatternFill('solid', fgColor='C6EFCE'); cell.font = Font(bold=True, color='006100')
    for col, width in {'A':29,'B':25,'C':14,'D':14,'E':12,'F':26}.items(): ws.column_dimensions[col].width = width
    notes = wb['Notes']; notes.column_dimensions['A'].width=24; notes.column_dimensions['B'].width=105
    for cell in notes[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
    wb.save(OUT_XLSX)
    # Refit selected pipeline on all labeled data to produce the deployment artifact.
    X_final = X_all.copy()
    if not use_count: X_final = X_final.drop(columns=['num_reported_accidents'])
    final_model = make_pipeline(X_final, winning_estimator)
    final_model.fit(X_final, y_all)
    artifact = {'model': final_model, 'algorithm': winning_name, 'accident_count_included': use_count,
                'feature_columns': X_final.columns.tolist(), 'target': TARGET, 'random_seed': SEED,
                'validation_metrics': winner, 'validation_selection': 'lowest MAE; RMSE tie-breaker',
                'trained_on_rows': len(X_final)}
    joblib.dump(artifact, OUT_MODEL, compress=3)
    included = {r['Algorithm']: r for r in rows if r['Accident Count Included']=='Yes'}
    excluded = {r['Algorithm']: r for r in rows if r['Accident Count Included']=='No'}
    diffs = {a: {'MAE_change_percent': (excluded[a]['MAE']-included[a]['MAE'])/excluded[a]['MAE']*100,
                 'RMSE_change_percent': (excluded[a]['RMSE']-included[a]['RMSE'])/excluded[a]['RMSE']*100}
             for a in included}
    best_by_config = {'with': min(included.values(),key=lambda x:(x['MAE'],x['RMSE'])),
                      'without': min(excluded.values(),key=lambda x:(x['MAE'],x['RMSE']))}
    delta_mae = (best_by_config['without']['MAE']-best_by_config['with']['MAE']) / best_by_config['without']['MAE'] * 100
    summary = {'winner': winning_name, 'accident_count_included': use_count, 'metrics': winner,
               'best_with_accident_count': best_by_config['with'], 'best_without_accident_count': best_by_config['without'],
               'best_configuration_mae_improvement_percent_with_count': delta_mae, 'per_algorithm_deltas_percent': diffs,
               'materiality_assessment': 'MAE difference >= 5% between best per-configuration models' if abs(delta_mae)>=5 else 'MAE difference < 5% between best per-configuration models; not material by this practical threshold',
               'test_labels_available': TARGET in pd.read_csv(TEST_FILE, nrows=0).columns}
    (ROOT/'experiment_summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print('\nWINNER:', winning_name, '| accident count included:', use_count)
    print(f"Validation MAE={winner['MAE']:.5f}, RMSE={winner['RMSE']:.5f}, R2={winner['R²']:.5f}, training_seconds={winner['Training Time (seconds)']:.1f}")
    print(f"Best-with vs best-without MAE improvement: {delta_mae:.2f}%")
    print('Test evaluation unavailable: test labels absent.' if TARGET not in pd.read_csv(TEST_FILE,nrows=0).columns else 'Test labels found.')
    print('Saved:', OUT_XLSX.name, OUT_MODEL.name)

if __name__ == '__main__': main()
